import io
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib import colors


def export_to_markdown(title: str, content: str) -> bytes:
    md = f"# {title}\n\n{content}\n"
    return md.encode("utf-8")


def export_to_docx(title: str, content: str) -> bytes:
    document = Document()
    document.add_heading(title, level=1)
    for line in content.split("\n"):
        if line.strip():
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.read()


def export_to_pdf(title: str, content: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleStyle", parent=styles["Title"], textColor=colors.HexColor("#6D28D9")
    )
    body_style = ParagraphStyle(
        "BodyStyle", parent=styles["BodyText"], fontSize=11, leading=16, spaceAfter=10
    )

    story = [Paragraph(title, title_style), Spacer(1, 16)]
    for line in content.split("\n"):
        if line.strip():
            safe_line = (
                line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            )
            story.append(Paragraph(safe_line, body_style))
    doc.build(story)
    buffer.seek(0)
    return buffer.read()
