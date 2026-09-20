import os
import fitz  # PyMuPDF
import docx
from pptx import Presentation
import markdown as md_lib
import re


def extract_text_from_pdf(path: str) -> str:
    text_parts = []
    with fitz.open(path) as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts).strip()


def extract_text_from_docx(path: str) -> str:
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()


def extract_text_from_pptx(path: str) -> str:
    prs = Presentation(path)
    parts = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        parts.append(f"--- Slide {slide_num} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs)
                    if line.strip():
                        parts.append(line)
    return "\n".join(parts).strip()


def extract_text_from_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read().strip()


def extract_text_from_markdown(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        raw = f.read()
    html = md_lib.markdown(raw)
    text = re.sub(r"<[^>]+>", " ", html)
    return text.strip()


def extract_text(path: str, file_type: str) -> str:
    file_type = file_type.lower().lstrip(".")
    if file_type == "pdf":
        return extract_text_from_pdf(path)
    if file_type == "docx":
        return extract_text_from_docx(path)
    if file_type == "pptx":
        return extract_text_from_pptx(path)
    if file_type in ("txt",):
        return extract_text_from_txt(path)
    if file_type in ("md", "markdown"):
        return extract_text_from_markdown(path)
    raise ValueError(f"Unsupported file type: {file_type}")


def get_file_extension(filename: str) -> str:
    return os.path.splitext(filename)[1].lower().lstrip(".")
